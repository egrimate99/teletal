import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import streamlit as st

from teletal_core import discover_menu_sections, discover_weeks, scrape_menu_rows


def build_xlsx(rows):
    fieldnames = list(rows[0].keys())
    numeric_cols = {
        "kcal_adag", "kj_adag", "zsír_g_adag", "zsír_telített_g_adag",
        "szénhidrát_g_adag", "cukor_g_adag", "rost_g_adag", "fehérje_g_adag",
        "só_g_adag", "kcal_100g", "kj_100g", "zsír_g_100g",
        "szénhidrát_g_100g", "fehérje_g_100g", "súly_g", "nap_szám", "hét", "év",
        "ár_ft",
    }
    wb = Workbook()
    ws = wb.active
    ev = rows[0].get("év", "")
    het = rows[0].get("hét", "")
    ws.title = f"{ev} {het}. hét"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E75B6")

    for col_idx, name in enumerate(fieldnames, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(fieldnames, 1):
            val = row[key]
            if key in numeric_cols and val != "":
                try:
                    val = float(val) if "." in str(val) else int(val)
                except ValueError:
                    pass
            ws.cell(row=row_idx, column=col_idx, value=val)

    for col_idx, name in enumerate(fieldnames, 1):
        max_len = max(len(str(name)), *(len(str(r[name])) for r in rows))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@st.cache_data(ttl=24 * 60 * 60, show_spinner=False)
def cached_menu_rows(week):
    rows, pages = scrape_menu_rows(
        week=week,
        delay_seconds=0,
    )
    return rows, pages, datetime.now().strftime("%Y-%m-%d %H:%M")


@st.cache_data(ttl=24 * 60 * 60)
def week_options():
    return discover_weeks()


@st.cache_data(ttl=24 * 60 * 60)
def menu_options(week):
    sections, _ = discover_menu_sections(week)
    return sections


def dataframe_from_rows(rows):
    df = pd.DataFrame(rows)
    numeric_cols = [
        "hét", "év", "nap_szám", "ár_ft", "súly_g", "kcal_adag", "kj_adag",
        "zsír_g_adag", "zsír_telített_g_adag", "szénhidrát_g_adag",
        "cukor_g_adag", "rost_g_adag", "fehérje_g_adag", "só_g_adag",
        "kcal_100g", "kj_100g", "zsír_g_100g", "szénhidrát_g_100g",
        "fehérje_g_100g",
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------

st.set_page_config(page_title="Teletál étlap", page_icon="🍽️", layout="wide")
st.title("🍽️ Teletál étlap")

try:
    weeks, active_week = week_options()
except Exception:
    weeks, active_week = [], None

with st.sidebar:
    if st.button("Frissítés most", use_container_width=True):
        cached_menu_rows.clear()
        menu_options.clear()
        week_options.clear()
        st.rerun()

if weeks:
    week = st.sidebar.selectbox(
        "Hét",
        weeks,
        index=weeks.index(active_week) if active_week in weeks else 0,
        format_func=lambda value: f"{value}. hét",
    )
else:
    week = None

if not week:
    st.error("Nem sikerült betölteni a választható heteket.")
    st.stop()

try:
    with st.spinner("Adatok betöltése..."):
        rows, pages, generated_at = cached_menu_rows(week)
        sections = menu_options(week)
except Exception as e:
    st.error(f"Hiba történt: {e}")
    st.stop()

if not rows:
    st.warning("Nem található letölthető étel.")
    st.stop()

df = dataframe_from_rows(rows)
ev = pages[0]["ev"] if pages else ""
het = pages[0]["het"] if pages else week

with st.sidebar:
    st.caption(f"Utolsó frissítés: {generated_at}")
    st.caption(f"{len(sections)} menü, {len(df)} sor/napi tétel")

    menu_values = sorted(df["menü"].dropna().unique())
    selected_menus = st.multiselect("Menü szűrő", menu_values, default=menu_values)

    day_values = list(df["nap"].dropna().unique())
    selected_days = st.multiselect("Nap szűrő", day_values, default=day_values)

    search = st.text_input("Keresés")

filtered = df[df["menü"].isin(selected_menus) & df["nap"].isin(selected_days)]
if search:
    search_mask = filtered.astype(str).apply(
        lambda col: col.str.contains(search, case=False, na=False)
    ).any(axis=1)
    filtered = filtered[search_mask]

col1, col2, col3, col4 = st.columns(4)
col1.metric("Menük", filtered["menü"].nunique())
col2.metric("Sor/napi tétel", len(filtered))
col3.metric("Kódok", filtered["kod"].nunique())
col4.metric("Hét", f"{ev}/{het}")

st.dataframe(
    filtered,
    use_container_width=True,
    hide_index=True,
    height=680,
)

if filtered.empty:
    st.info("Nincs találat a jelenlegi szűrőkkel.")
else:
    xlsx_buf = build_xlsx(filtered.to_dict("records"))
    st.download_button(
        label="Excel export",
        data=xlsx_buf,
        file_name=f"teletal_{ev}_{het}het.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
