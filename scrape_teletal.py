"""
Teletál menu scraper.

Fetches the selected week's Teletál menu sections with nutrition data,
then writes CSV and XLSX exports.
"""

import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from teletal_core import MENU_SOURCES, scrape_menu_rows


NUMERIC_COLS = {
    "kcal_adag",
    "kj_adag",
    "zsír_g_adag",
    "zsír_telített_g_adag",
    "szénhidrát_g_adag",
    "cukor_g_adag",
    "rost_g_adag",
    "fehérje_g_adag",
    "só_g_adag",
    "kcal_100g",
    "kj_100g",
    "zsír_g_100g",
    "szénhidrát_g_100g",
    "fehérje_g_100g",
    "súly_g",
    "nap_szám",
    "hét",
    "év",
    "ár_ft",
}


def numeric_value(key: str, value):
    if key not in NUMERIC_COLS or value == "":
        return value
    try:
        return float(value) if "." in str(value) else int(value)
    except ValueError:
        return value


def write_csv(rows: list[dict], filename: str) -> None:
    fieldnames = list(rows[0].keys())
    with open(filename, "w", newline="", encoding="utf-8-sig") as csvfile:
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows: list[dict], filename: str) -> None:
    fieldnames = list(rows[0].keys())
    wb = Workbook()
    ws = wb.active
    ws.title = f"{rows[0].get('év', '')} {rows[0].get('hét', '')}. hét"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E75B6")
    center = Alignment(horizontal="center")

    for col_idx, name in enumerate(fieldnames, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(fieldnames, 1):
            ws.cell(row=row_idx, column=col_idx, value=numeric_value(key, row[key]))

    for col_idx, name in enumerate(fieldnames, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(name)), *(len(str(row[name])) for row in rows))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 45)

    ws.freeze_panes = "A2"
    wb.save(filename)


def scrape_menu(week: str | None = None) -> None:
    print("Fetching menu pages:")
    for source in MENU_SOURCES:
        print(f"  - {source.label}")

    def show_progress(index: int, total: int, item: dict) -> None:
        print(
            f"  [{index}/{total}] {item['menu_title'][:28]:<28} "
            f"kod={item['kod']:<6} nap={item['nap']} ({item['nap_nev']})"
        )

    rows, pages = scrape_menu_rows(
        week=week,
        progress_callback=show_progress,
        delay_seconds=0.1,
    )
    if not rows:
        print("No data found.")
        return

    csv_file = "teletal_menu.csv"
    xlsx_file = "teletal_menu.xlsx"
    write_csv(rows, csv_file)
    write_xlsx(rows, xlsx_file)

    print("Done!")
    print(f"  Rows: {len(rows)}")
    for page in pages:
        print(f"  {page['source']}: {page['item_count']} items from {page['url']}")
    print(f"  Wrote {csv_file}")
    print(f"  Wrote {xlsx_file}")


if __name__ == "__main__":
    scrape_menu()
